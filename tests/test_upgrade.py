import copy
import importlib.util
import json
import math
from pathlib import Path
import random
import sys
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from unittest.mock import patch, MagicMock

import pytest
from portfolio_config import validate_config
from portfolio_llm import parse_config, generate_portfolio_config, run_config_tool
from portfolio_budget import allocate_shares
from portfolio_cli import main, run_portfolio

BASE = {'total_value': 38000, 'num_picks': 10}

@pytest.mark.parametrize('field,value', [
    ('total_value', 0), ('total_value', True), ('total_value', float('nan')),
    ('flat_fee', -1), ('num_picks', 0), ('num_picks', True), ('num_picks', 101),
    ('top_n_sectors', 12), ('filter_pe_ratio', 'true'), ('performance_period', '6m'),
    ('historical', []), ('currency', 'usd'), ('cache', 'evil'),
    ('output_filename', '../evil.xlsx'), ('output_filename', '/tmp/evil.xlsx'),
    ('output_filename', 'x.csv'), ('picks', {'MSFT': 101}), ('picks', {'MSFT': -1}),
    ('picks', {'MSFT': float('inf')}), ('picks', {'bad/name': 10}),
    ('picks', ['MSFT', 'MSFT']), ('picks', ['MSFT', 2]), ('banned_sectors', 'Energy'),
    ('unknown', 1), ('picks', {'MSFT': 60, 'GOOG': 60})])
def test_invalid_config(field, value):
    with pytest.raises(ValueError):
        validate_config(dict(BASE, **{field: value}))

def test_defaults_and_lists():
    original = dict(BASE, picks=['MSFT', 'GOOG'])
    assert validate_config(original)['picks'] == {'MSFT': 50, 'GOOG': 50}
    assert original['picks'] == ['MSFT', 'GOOG']

def test_parse_fence_and_reject_duplicate():
    assert parse_config('```json\n' + json.dumps(BASE) + '\n```')['total_value'] == 38000
    with pytest.raises(ValueError):
        parse_config('{"total_value":1000,"num_picks":1,"num_picks":2}')

@pytest.mark.parametrize('text', ['nope', '[]', '', '{"total_value": NaN,"num_picks":1}', 'a' * 65537])
def test_bad_model_json(text):
    with pytest.raises(ValueError): parse_config(text)

@pytest.mark.parametrize('arguments', [{}, {'request':'x','base_url':'http://evil'}, [], {'request':''}])
def test_tool_rejects_arguments(arguments):
    with pytest.raises(ValueError): run_config_tool(arguments)


def test_local_qwen_http_contract():
    seen = {}
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            seen['path'] = self.path
            seen['body'] = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
            self.send_response(200); self.end_headers()
            self.wfile.write(json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(BASE)}}]}).encode())
        def log_message(self, *args): pass
    server = HTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    try:
        result = generate_portfolio_config('Research portfolio', provider='qwen', base_url=f'http://127.0.0.1:{server.server_port}/v1')
        assert result['total_value'] == 38000
        assert seen['path'] == '/v1/chat/completions'
        assert 'Qwen3.6' in seen['body']['model']
        assert seen['body']['stream'] is False
    finally: server.shutdown(); thread.join(); server.server_close()

@pytest.mark.parametrize('body', [b'{}', b'bad', json.dumps({'choices':[{'finish_reason':'length','message':{'content':json.dumps(BASE)}}]}).encode(), b'x' * 131073])
def test_bad_envelope(body):
    response = MagicMock(); response.__enter__.return_value.read.return_value = body
    opener = MagicMock(); opener.open.return_value = response
    with patch('urllib.request.build_opener', return_value=opener), pytest.raises(ValueError):
        generate_portfolio_config('x', provider='qwen')


def test_network_failure_no_cloud_fallback():
    with patch('urllib.request.build_opener') as factory:
        factory.return_value.open.side_effect = TimeoutError('test timeout')
        with pytest.raises(TimeoutError): generate_portfolio_config('x', provider='qwen')
        assert factory.return_value.open.call_count == 1

@pytest.mark.parametrize('url', ['file:///tmp/a','http://user:password@example.com/v1','http://example.com/v1?token=a'])
def test_endpoint_validation(url):
    with pytest.raises(ValueError): generate_portfolio_config('x', provider='qwen', base_url=url)


def test_redirect_is_not_followed():
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            self.send_response(307); self.send_header('Location', '/secret-target'); self.end_headers()
        def log_message(self,*args): pass
    server=HTTPServer(('127.0.0.1',0),Handler)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        from urllib.error import HTTPError
        with pytest.raises(HTTPError):
            generate_portfolio_config('x',provider='qwen',base_url=f'http://127.0.0.1:{server.server_port}/v1')
    finally: server.shutdown();thread.join();server.server_close()


def test_budget_fee_regression():
    counts, investments = allocate_shares([100, 100], 200, 9.99)
    assert sum(investments) <= 200
    assert counts == [0, 0]  # A share plus fee would exceed each 50% cap.

@pytest.mark.parametrize('prices,budget,fee', [([],100,1),([0],100,1),([-1],100,1),([math.nan],100,1),([10],0,1),([10],100,-1)])
def test_invalid_budget(prices,budget,fee):
    with pytest.raises(ValueError): allocate_shares(prices,budget,fee)


def test_randomized_budget_invariant():
    rng = random.Random(20261018)
    for _ in range(500):
        n=rng.randint(1,20);fee=rng.uniform(0,10);budget=rng.uniform(n*fee+1,100000)
        counts, investments=allocate_shares([rng.uniform(.01,10000) for _ in range(n)],budget,fee)
        assert all(type(c) is int and c >= 0 for c in counts)
        assert sum(investments) <= budget+1e-8


def test_large_budget_is_batched():
    counts,_ = allocate_shares([.01, .02], 10**12, 1)
    assert sum(counts) > 10**12


def test_manual_cli_has_no_model_or_market_side_effects(tmp_path):
    path=tmp_path/'config.json';path.write_text(json.dumps(BASE))
    with patch('portfolio_llm.generate_portfolio_config',side_effect=AssertionError('LLM called')):
        assert main(['--config',str(path),'--validate-only'])['total_value']==38000


def test_real_engine_and_excel_without_market(tmp_path):
    pd=pytest.importorskip('pandas')
    engine=pytest.importorskip('make_fast_portfolio')
    selected=pd.DataFrame({'Ticker':['MSFT','GOOG'],'Share_Price':[100.,80.], 'Percentage':[50.,50.]})
    with patch.object(engine,'get_exchange_rate',return_value=1), patch.object(engine,'select_top_stocks',return_value=(selected,{'Tech':1.2})):
        summary=run_portfolio(dict(BASE,num_picks=2),output_dir=tmp_path,forecasts=False)
    from openpyxl import load_workbook
    workbook=load_workbook(summary['output_file'])
    assert workbook.sheetnames == ['Top Stocks','Sector Performance']
    assert summary['cash_remaining']>=0
    assert Path(summary['output_file']).with_suffix('.summary.json').is_file()


def load_dag_with_stub():
    """Exercise pure DAG adapter without pretending this is Airflow scheduler validation."""
    import types
    sdk=types.ModuleType('airflow.sdk')
    sdk.dag=lambda **kw: lambda f: f
    sdk.task=lambda **kw: lambda f: lambda: None
    sdk.get_current_context=lambda: {}
    spec=importlib.util.spec_from_file_location('dag_under_test',Path(__file__).parents[1]/'dags/stock_portfolio_research.py')
    module=importlib.util.module_from_spec(spec)
    with patch.dict(sys.modules,{'airflow':types.ModuleType('airflow'),'airflow.sdk':sdk}):
        spec.loader.exec_module(module)
    return module


def test_dag_adapter_uses_pipeline_and_isolates_paths(tmp_path,monkeypatch):
    module=load_dag_with_stub();monkeypatch.setenv('PORTFOLIO_OUTPUT_ROOT',str(tmp_path))
    with patch('portfolio_cli.run_portfolio',return_value={'ok':True}) as run:
        assert module.execute_run({'config':BASE,'forecasts':False},'manual/../../escape') == {'ok':True}
        output=run.call_args.kwargs['output_dir']
        assert output.parent==tmp_path and len(output.name)==24
        assert run.call_args.kwargs['plots'] is False

@pytest.mark.parametrize('conf',[{'config':BASE,'request':'x'},{'base_url':'http://evil'},{'forecasts':'false'},[]])
def test_dag_rejects_unsafe_conf(conf):
    with pytest.raises(ValueError): load_dag_with_stub().execute_run(conf,'run')


def test_gemini_compatibility_is_lazy_and_validated(monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY','synthetic-not-a-secret')
    from google import genai
    with patch.object(genai,'Client') as client:
        client.return_value.models.generate_content.return_value.text=json.dumps(BASE)
        from cgemini1206exp import generate_portfolio_config as legacy
        assert legacy('A research portfolio')['num_picks']==10
        assert client.call_args.kwargs['http_options']['timeout']==60000


def test_legacy_selection_pipeline_with_synthetic_fundamentals(tmp_path):
    import pandas as pd
    import make_fast_portfolio as engine
    row={'Ticker':'MSFT','Sector':'Tech','EPS':2.,'Revenue':100.,'MarketCap':1000.,
         'PE_Ratio':10.,'Price_to_Sales':2.,'Price_to_Book':2.,'EBITDA':30.,
         'Profit_Margins':.2,'Share_Price':100.,'Quick_Ratio':1.5}
    config=validate_config(dict(BASE,num_picks=1,picks={'MSFT':100}))
    table=pd.DataFrame({'Symbol':['MSFT'],'Security':['Microsoft'],'GICS Sector':['Tech']})
    with patch.object(engine,'get_sp500_tickers',return_value=(['MSFT'],table)), patch.object(engine,'get_sector_performance',return_value={'Tech':2}), patch.object(engine,'analyze_fundamentals',return_value=row):
        selected,sectors=engine.select_top_stocks(config,stock_picks=config['picks'])
    assert selected['Ticker'].tolist()==['MSFT']


def test_genuine_prophet_forecast_export_with_synthetic_history(tmp_path):
    import pandas as pd
    import make_fast_portfolio as engine
    from openpyxl import load_workbook
    history=pd.DataFrame({'Close':[100+i*.2 for i in range(40)]},index=pd.date_range('2026-01-01',periods=40,name='Date'))
    portfolio=pd.DataFrame({'Ticker':['MSFT'],'Share_Price':[108.]})
    path=tmp_path/'forecast.xlsx'
    with patch.object(engine,'fetch_stock_forecast',return_value={'MSFT':history}):
        engine.save_forecast_to_excel_predict(portfolio,{'Tech':1.2},str(path),validate_config(BASE))
    assert 'MSFT_Forecast' in load_workbook(path).sheetnames


def test_review_regression_partial_allocation_preserves_cash():
    counts,investments=allocate_shares([100,100],1000,0,[25,25])
    assert counts==[2,2] and sum(investments)==400


def test_review_regression_zero_weight_never_funded():
    counts,_=allocate_shares([600,100],1000,0,[100,0])
    assert counts==[1,0]


def test_review_regression_epsilon_weight_never_overspends():
    counts,investments=allocate_shares([1,1],10**12,0,[50.000000004,50.000000004])
    assert sum(investments)<=10**12


def test_review_regression_fees_only_charged_to_funded_positions():
    counts,investments=allocate_shares([1000,10],100,10,[0,100])
    assert counts==[0,9] and sum(investments)==100
